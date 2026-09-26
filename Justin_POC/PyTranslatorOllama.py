#!/usr/bin/env python3

"""
Translate a PowerPoint presentation from English to Spanish
using a local Ollama instance running Gemma.

Requirements:
    pip install python-pptx requests

Usage:
    python translate_ppt.py input.pptx output.pptx
"""

import argparse
import json
import requests
from pptx import Presentation

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL = "gemma4:e2b"


def translate_text(text: str) -> str:
    """Translate English text to Spanish using Ollama."""

    text = text.strip()

    if not text:
        return text

    prompt = f"""
Translate the following English text into natural Spanish.

Rules:
- Return ONLY the translated text.
- Preserve technical terminology when appropriate.
- Preserve line breaks.
- Do not add explanations.
- Do not add quotation marks.

Text:
{text}
"""

    response = requests.post(
        OLLAMA_URL,
        json={
            "model": MODEL,
            "prompt": prompt,
            "stream": False,
            "thinking": False,
            "options": {
                "temperature": 0.1
            }
        },
        timeout=300,
    )

    response.raise_for_status()

    result = response.json()
    return result.get("response", "").strip()


def translate_text_frame(text_frame):
    """Translate all text in a PowerPoint text frame."""

    original_text = text_frame.text

    if not original_text.strip():
        return

    try:
        translated = translate_text(original_text)
        text_frame.text = translated

        print("-" * 60)
        print("EN:", original_text[:120])
        print("ES:", translated[:120])

    except Exception as e:
        print(f"Translation failed: {e}")


def process_shape(shape):
    """Recursively process PowerPoint shapes."""

    # Text boxes, placeholders, titles, etc.
    if hasattr(shape, "text_frame") and shape.has_text_frame:
        translate_text_frame(shape.text_frame)

    # Tables
    if shape.has_table:
        table = shape.table
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    try:
                        translated = translate_text(cell.text)
                        cell.text = translated
                    except Exception as e:
                        print(f"Table translation failed: {e}")

    # Groups
    if shape.shape_type == 6:  # GROUP
        for subshape in shape.shapes:
            process_shape(subshape)


def translate_presentation(input_file, output_file):
    """Translate an entire presentation."""

    prs = Presentation(input_file)

    total_slides = len(prs.slides)

    for slide_num, slide in enumerate(prs.slides, start=1):
        print(f"\nProcessing slide {slide_num}/{total_slides}")

        for shape in slide.shapes:
            process_shape(shape)

    prs.save(output_file)

    print(f"\nTranslated presentation saved to:")
    print(output_file)


def main():
    parser = argparse.ArgumentParser(
        description="Translate PowerPoint from English to Spanish using Ollama"
    )

    parser.add_argument("input", help="Input PPTX file")
    parser.add_argument("output", help="Output PPTX file")

    args = parser.parse_args()

    translate_presentation(args.input, args.output)


if __name__ == "__main__":
    main()