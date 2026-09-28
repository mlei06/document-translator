# /// script
# requires-python = ">=3.14"
# dependencies = ["python-pptx==1.0.2", "python-docx==1.2.0", "lxml==6.1.3"]
# ///
"""PPTX/DOCX writer experiment for P2.0, not a production adapter.

Run from the repository root after make_fixtures.ps1:

    uv run docs/experiments/ooxml-roundtrip/spike.py

For deck.pptx and report.docx (authored by native PowerPoint/Word), compare:
- library: python-pptx / python-docx load, prefix every run reachable through the library's
  object model with "T:", save;
- targeted: copy every ZIP entry byte-for-byte except the text-bearing XML parts, in which every
  a:t / w:t text node is prefixed with "T:" (lxml, namespaces preserved), then serialize only
  those parts.

Records per writer: ZIP entries whose uncompressed payload changed, entries missing from the output,
text nodes in the package that were not changed (coverage), and the package text inventory.
Writes outputs and results.json under data/experiments/ooxml-roundtrip/. Native-open checks are run
separately with scripts/native_office_check.ps1 on the outputs.
"""

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path

import docx
import pptx
from lxml import etree

ROOT = Path("data/experiments/ooxml-roundtrip")
FIXTURES = ROOT / "fixtures"
OUTPUTS = ROOT / "outputs"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PREFIX = "T:"
PARSER = etree.XMLParser(resolve_entities=False, no_network=True, remove_blank_text=False)

PPTX_TEXT_PART = re.compile(r"ppt/(slides/slide|notesSlides/notesSlide)\d+\.xml$")
DOCX_TEXT_PART = re.compile(r"word/(document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml$")


def entries(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def text_nodes(data: dict[str, bytes]) -> dict[str, list[str]]:
    """Every a:t / w:t text node in any XML part, by part name."""
    found: dict[str, list[str]] = {}
    for name, payload in data.items():
        if not name.endswith(".xml"):
            continue
        root = etree.fromstring(payload, PARSER)
        texts = [t.text or "" for t in root.iter(f"{{{A}}}t", f"{{{W}}}t")]
        if texts:
            found[name] = texts
    return found


def targeted(source: Path, target: Path, pattern: re.Pattern[str]) -> list[str]:
    changed: list[str] = []
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w") as dst:
        for info in src.infolist():
            payload = src.read(info.filename)
            if pattern.search(info.filename):
                root = etree.fromstring(payload, PARSER)
                nodes = [t for t in root.iter(f"{{{A}}}t", f"{{{W}}}t") if t.text]
                for node in nodes:
                    node.text = PREFIX + (node.text or "")
                if nodes:
                    payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    changed.append(info.filename)
            dst.writestr(info, payload)
    return changed


def library_pptx(source: Path, target: Path) -> None:
    pres = pptx.Presentation(str(source))

    def shapes(collection):  # type: ignore[no-untyped-def]
        for shape in collection:
            if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                yield from shapes(shape.shapes)
            else:
                yield shape

    def frames(slide):  # type: ignore[no-untyped-def]
        for shape in shapes(slide.shapes):
            if shape.has_text_frame:
                yield shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame

    for slide in pres.slides:
        all_frames = list(frames(slide))
        if slide.has_notes_slide:
            all_frames.append(slide.notes_slide.notes_text_frame)
        for frame in all_frames:
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    if run.text:
                        run.text = PREFIX + run.text
    pres.save(str(target))


def library_docx(source: Path, target: Path) -> None:
    document = docx.Document(str(source))

    def paragraphs(container):  # type: ignore[no-untyped-def]
        yield from container.paragraphs
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)

    blocks = [document]
    for section in document.sections:
        blocks += [section.header, section.footer]
    for block in blocks:
        for paragraph in paragraphs(block):
            for run in paragraph.runs:
                if run.text:
                    run.text = PREFIX + run.text
    document.save(str(target))


def compare(source: Path, output: Path) -> dict[str, object]:
    before, after = entries(source), entries(output)
    changed = sorted(n for n in before if n in after and before[n] != after[n])
    missing = sorted(set(before) - set(after))
    added = sorted(set(after) - set(before))
    src_text, out_text = text_nodes(before), text_nodes(after)
    untranslated = {
        part: [t for t in texts if t and not t.startswith(PREFIX)]
        for part, texts in out_text.items()
    }
    untranslated = {p: t for p, t in untranslated.items() if t}
    return {
        "changed_entries": changed,
        "missing_entries": missing,
        "added_entries": added,
        "text_nodes_source": sum(len(t) for t in src_text.values()),
        "untranslated_text_nodes": untranslated,
        "sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
    }


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")  # pyright: ignore[reportAttributeAccessIssue]
    OUTPUTS.mkdir(parents=True, exist_ok=True)
    results: dict[str, object] = {"versions": {"python-pptx": pptx.__version__, "lxml": etree.__version__}}
    cases = [
        ("deck.pptx", library_pptx, PPTX_TEXT_PART),
        ("report.docx", library_docx, DOCX_TEXT_PART),
    ]
    for name, library, pattern in cases:
        source = FIXTURES / name
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        lib_out = OUTPUTS / f"library-{name}"
        tgt_out = OUTPUTS / f"targeted-{name}"
        library(source, lib_out)
        edited = targeted(source, tgt_out, pattern)
        assert hashlib.sha256(source.read_bytes()).hexdigest() == digest, "input modified"
        results[name] = {
            "source_parts_with_text": sorted(text_nodes(entries(source))),
            "library": compare(source, lib_out),
            "targeted": compare(source, tgt_out) | {"edited_parts": edited},
        }
    (ROOT / "results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
