# /// script
# requires-python = ">=3.14"
# dependencies = ["openpyxl==3.1.5", "xlsxwriter==3.2.5", "pillow==12.1.1", "lxml==6.0.2"]
# ///
"""Synthetic preservation experiment, not a production XLSX adapter.

Run: uv run docs/experiments/xlsx-roundtrip/spike.py
Writes disposable workbooks and results under data/experiments/xlsx-roundtrip/.
No translation server, confidential input, Excel automation, or project dependency changes.
"""

import copy
import hashlib
import io
import json
import platform
import sys
import warnings
import zipfile
from datetime import datetime
from importlib.metadata import version
from pathlib import Path

import xlsxwriter
from lxml import etree
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from PIL import Image

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
S = "{" + NS["s"] + "}"
OUT = Path("data/experiments/xlsx-roundtrip")
SOURCE_NAME = "销售数据"
TARGET_NAME = "Sales"
PREFIX = "T:"


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def parts(path):
    with zipfile.ZipFile(path) as archive:
        require(archive.testzip() is None, f"bad ZIP: {path}")
        return {name: archive.read(name) for name in archive.namelist()}


def parse(data):
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))


def write_parts(source, target, replacements):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, replacements.get(item.filename, src.read(item.filename)))


def fixture(path):
    with xlsxwriter.Workbook(path) as book:
        sheet = book.add_worksheet(SOURCE_NAME)
        notes = book.add_worksheet("说明")
        bold = book.add_format({"bold": True, "font_color": "red"})
        date_format = book.add_format({"num_format": "yyyy-mm-dd"})
        sheet.set_column("A:A", 22)
        sheet.set_row(0, 24)
        sheet.freeze_panes(1, 0)
        sheet.write("A1", "地区", bold)
        sheet.write("B1", "销售额")
        sheet.write("C1", "日期")
        for row, (name, amount) in enumerate((("华东", 120), ("华南", 95), ("华北", 80)), 1):
            sheet.write(row, 0, name)
            sheet.write(row, 1, amount)
            sheet.write_datetime(row, 2, datetime(2026, 9, 1), date_format)
        sheet.write_formula("D2", "=SUM(B2:B4)", None, 295)
        sheet.write_rich_string("E1", "普通文字", bold, "加粗文字")
        sheet.write_comment("A2", "请核对", {"author": "测试"})
        sheet.data_validation("F2", {"validate": "list", "source": ["是", "否"]})
        sheet.conditional_format(
            "B2:B4", {"type": "cell", "criteria": ">", "value": 100, "format": bold}
        )
        sheet.merge_range("A7:C7", "合并标题", bold)
        sheet.write_boolean("G2", True)
        sheet.write_string("G3", "=literal-not-a-formula")
        sheet.write_url("A6", "https://example.invalid/manual", string="手册")
        sheet.insert_textbox("H12", "请保留此形状")
        chart = book.add_chart({"type": "column"})
        chart.add_series(
            {"values": [SOURCE_NAME, 1, 1, 3, 1], "categories": [SOURCE_NAME, 1, 0, 3, 0]}
        )
        sheet.insert_chart("H2", chart)
        image = io.BytesIO()
        Image.new("RGB", (40, 20), "blue").save(image, format="PNG")
        sheet.insert_image("H20", "blue.png", {"image_data": image})
        notes.write("A1", "内部参考")
        notes.write_formula("B1", f"='{SOURCE_NAME}'!B2", None, 120)
        notes.write_url("A2", f"internal:'{SOURCE_NAME}'!A1", string="返回")
        book.define_name("SalesAmount", f"='{SOURCE_NAME}'!$B$2")


def inline_fixture(source, target):
    original = parts(source)
    strings = parse(original["xl/sharedStrings.xml"]).findall(f"{S}si")
    updates = {}
    for name, data in original.items():
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = parse(data)
        for cell in root.xpath("//s:c[@t='s']", namespaces=NS):
            value = cell.find(f"{S}v")
            text = copy.deepcopy(strings[int(value.text)])
            text.tag = f"{S}is"
            cell.remove(value)
            cell.set("t", "inlineStr")
            cell.append(text)
        updates[name] = etree.tostring(root)
    # Keeping an unused shared-string table is valid; it also tests unreferenced part preservation.
    write_parts(source, target, updates)


def translate_library(source, target, rename=False):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        book = load_workbook(source, rich_text=True)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "f":
                        continue
                    if isinstance(cell.value, CellRichText):
                        cell.value = CellRichText(
                            [
                                TextBlock(copy.copy(run.font), PREFIX + run.text)
                                if isinstance(run, TextBlock)
                                else PREFIX + run
                                for run in cell.value
                            ]
                        )
                    elif isinstance(cell.value, str):
                        cell.value = PREFIX + cell.value
        if rename:
            book[SOURCE_NAME].title = TARGET_NAME
        book.save(target)
        book.close()
    return [str(item.message) for item in caught]


def translate_xml(source, target, rename=False):
    updates = {}
    changed_text_nodes = 0
    for name, data in parts(source).items():
        if name == "xl/sharedStrings.xml":
            root = parse(data)
            nodes = root.xpath("/s:sst/s:si/s:t | /s:sst/s:si/s:r/s:t", namespaces=NS)
        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            root = parse(data)
            nodes = root.xpath(
                "//s:c[@t='inlineStr']/s:is/s:t | //s:c[@t='inlineStr']/s:is/s:r/s:t", namespaces=NS
            )
        elif name == "xl/workbook.xml" and rename:
            root = parse(data)
            nodes = []
            root.xpath("//s:sheet[@name=$name]", namespaces=NS, name=SOURCE_NAME)[0].set(
                "name", TARGET_NAME
            )
            updates[name] = etree.tostring(root)
        else:
            continue
        for node in nodes:
            node.text = PREFIX + (node.text or "")
            changed_text_nodes += 1
        if nodes:
            updates[name] = etree.tostring(root)
    write_parts(source, target, updates)
    require(changed_text_nodes > 0, "XML branch was a no-op")


def snapshot(path):
    archive = parts(path)
    book = load_workbook(path, rich_text=True)
    cached = load_workbook(path, data_only=True)
    sheet, notes = book.worksheets
    drawing = parse(archive["xl/drawings/drawing1.xml"])
    result = {
        "sheet_names": book.sheetnames,
        "plain_text": sheet["A1"].value,
        "rich_text": str(sheet["E1"].value),
        "rich_runs": repr(sheet["E1"].value),
        "bold_run": isinstance(sheet["E1"].value, CellRichText)
        and any(isinstance(run, TextBlock) and run.font.b for run in sheet["E1"].value),
        "local_formula": sheet["D2"].value,
        "cross_sheet_formula": notes["B1"].value,
        "formula_cache": cached.worksheets[0]["D2"].value,
        "cross_sheet_cache": cached.worksheets[1]["B1"].value,
        "number": sheet["B2"].value,
        "date": str(sheet["C2"].value),
        "boolean": sheet["G2"].value,
        "literal_equals": sheet["G3"].value,
        "comment": sheet["A2"].comment.text,
        "validation": str(sheet.data_validations),
        "conditional_rules": len(sheet.conditional_formatting),
        "merged": str(sheet.merged_cells),
        "cell_style": str(sheet["A1"]._style),
        "width": sheet.column_dimensions["A"].width,
        "height": sheet.row_dimensions[1].height,
        "freeze": sheet.freeze_panes,
        "chart_count": len(sheet._charts),
        "image_count": len(sheet._images),
        "shape_count": len(drawing.xpath("//*[local-name()='sp']")),
        "defined_name": book.defined_names["SalesAmount"].attr_text,
        "internal_link": notes["A2"].hyperlink.location,
        "chart_references": parse(archive["xl/charts/chart1.xml"]).xpath(
            "//*[local-name()='f']/text()"
        ),
    }
    book.close()
    cached.close()
    return result


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    OUT.mkdir(parents=True, exist_ok=True)
    fixture(OUT / "shared.xlsx")
    inline_fixture(OUT / "shared.xlsx", OUT / "inline.xlsx")
    results = {
        "python": platform.python_version(),
        "versions": {name: version(name) for name in ("openpyxl", "xlsxwriter", "pillow", "lxml")},
        "cases": {},
    }
    for storage in ("shared", "inline"):
        source = OUT / f"{storage}.xlsx"
        before_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        original = parts(source)
        baseline = snapshot(source)
        for writer in ("library", "xml"):
            for rename in (False, True):
                case = f"{storage}-{writer}-{'rename' if rename else 'keep-names'}"
                target = OUT / f"{case}.xlsx"
                messages = (
                    translate_library(source, target, rename)
                    if writer == "library"
                    else translate_xml(source, target, rename)
                )
                after = parts(target)
                observed = snapshot(target)
                require(observed["plain_text"] == PREFIX + baseline["plain_text"], case)
                require(
                    observed["rich_text"] == "T:普通文字T:加粗文字" and observed["bold_run"], case
                )
                invariants = (
                    "local_formula",
                    "cross_sheet_formula",
                    "number",
                    "date",
                    "boolean",
                    "comment",
                    "validation",
                    "conditional_rules",
                    "merged",
                    "cell_style",
                    "width",
                    "height",
                    "freeze",
                    "chart_count",
                    "image_count",
                )
                require(
                    all(observed[key] == baseline[key] for key in invariants),
                    f"semantic invariant failed: {case}",
                )
                for key in ("defined_name", "internal_link", "chart_references"):
                    require(observed[key] == baseline[key], f"reference changed: {case}: {key}")
                if rename:
                    require(SOURCE_NAME not in observed["sheet_names"], case)
                    for key in ("cross_sheet_formula", "defined_name", "internal_link"):
                        require(SOURCE_NAME in observed[key], f"expected stale reference: {case}")
                    require(all(SOURCE_NAME in ref for ref in observed["chart_references"]), case)
                changed = sorted(
                    name for name in original.keys() & after.keys() if original[name] != after[name]
                )
                lost = sorted(original.keys() - after.keys())
                if writer == "xml":
                    allowed = {
                        "xl/sharedStrings.xml",
                        "xl/worksheets/sheet1.xml",
                        "xl/worksheets/sheet2.xml",
                    }
                    if rename:
                        allowed.add("xl/workbook.xml")
                    require(not lost and not (after.keys() - original.keys()), case)
                    require(set(changed) <= allowed, case)
                    require(observed["formula_cache"] == 295 and observed["shape_count"] == 1, case)
                results["cases"][case] = {
                    "lost_parts": lost,
                    "changed_parts": changed,
                    "warnings": messages or [],
                    "observed": observed,
                    "stale_sheet_reference": rename
                    and SOURCE_NAME in observed["cross_sheet_formula"],
                }
        require(hashlib.sha256(source.read_bytes()).hexdigest() == before_hash, "input modified")
    (OUT / "results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    for name, result in results["cases"].items():
        state = result["observed"]
        print(
            f"{name}: changed={len(result['changed_parts'])}, lost={result['lost_parts']}, "
            f"shapes={state['shape_count']}, cache={state['formula_cache']}, "
            f"stale_reference={result['stale_sheet_reference']}"
        )
    print(f"PASS: 8 cases; results at {OUT / 'results.json'}")


if __name__ == "__main__":
    main()
