import docx
from docx import Document
import sys

# Force UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

doc = Document(r"C:\Users\jringuette\repos\PyTranslator\samples\how to replace TC.docx")

print(f"Document loaded: {len(doc.paragraphs)} paragraphs")
print(f"Tables: {len(doc.tables)}")
print("=" * 60)

print("\n=== PARAGRAPHS ===")
for i, para in enumerate(doc.paragraphs):
    print(f"\n--- Paragraph {i} ---")
    print(f"  Style: {para.style.name}")
    print(f"  Text: {repr(para.text[:200])}")
    print(f"  Runs: {len(para.runs)}")
    for j, run in enumerate(para.runs):
        color = run.font.color.rgb if run.font.color and run.font.color.rgb else 'None'
        print(f"    Run {j}: text={repr(run.text[:100])}, bold={run.bold}, italic={run.italic}, underline={run.underline}, font={run.font.name}, size={run.font.size}, color={color}")

print("\n" + "=" * 60)
print("\n=== TABLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Table {t_idx} ---")
    print(f"  Rows: {len(table.rows)}, Columns: {len(table.columns)}")
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            print(f"  Cell [{r_idx}][{c_idx}]: {repr(cell.text[:100])}")
            for p_idx, para in enumerate(cell.paragraphs):
                print(f"    Para {p_idx}: {repr(para.text[:100])}")
                for r_idx, run in enumerate(para.runs):
                    print(f"      Run {r_idx}: {repr(run.text[:80])}, bold={run.bold}, italic={run.italic}")

print("\n" + "=" * 60)
print("\n=== DOCUMENT STYLES ===")
for style in doc.styles:
    if style.type == 1:  # paragraph style
        print(f"  Paragraph Style: {style.name}")

print("\n=== RUNS WITH FORMATTING (bold/italic/underline) ===")
for i, para in enumerate(doc.paragraphs):
    for j, run in enumerate(para.runs):
        if run.bold or run.italic or run.underline or (run.font.color and run.font.color.rgb):
            color = run.font.color.rgb if run.font.color and run.font.color.rgb else 'None'
            print(f"  Para {i}, Run {j}: {repr(run.text[:80])} | bold={run.bold}, italic={run.italic}, underline={run.underline}, color={color}, font={run.font.name}, size={run.font.size}")

print("\n=== TABLE CELL RUNS WITH FORMATTING ===")
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, para in enumerate(cell.paragraphs):
                for r_idx, run in enumerate(para.runs):
                    if run.bold or run.italic or run.underline or (run.font.color and run.font.color.rgb):
                        color = run.font.color.rgb if run.font.color and run.font.color.rgb else 'None'
                        print(f"  Table {t_idx}, Cell[{r_idx}][{c_idx}], Para {p_idx}, Run {r_idx}: {repr(run.text[:80])} | bold={run.bold}, italic={run.italic}, underline={run.underline}, color={color}")

print("\n=== TABLE STYLES ===")
for t_idx, table in enumerate(doc.tables):
    print(f"  Table {t_idx} style: {table.style}")

print("\n=== DOCUMENT PROPERTIES ===")
core_props = doc.core_properties
print(f"  Title: {core_props.title}")
print(f"  Author: {core_props.author}")
print(f"  Created: {core_props.created}")
print(f"  Modified: {core_props.modified}")

print("\n=== DONE ===")