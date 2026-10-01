"""Write the synthetic sample documents behind "Feed me sample files" (P6 audit row 13).

All text is invented for this purpose; nothing comes from real documents. Run from the repository
root with ``uv run --with python-pptx --with python-docx --with openpyxl python
apps/web/scripts/make_samples.py`` (one-off tools, not project dependencies); the files land in
``apps/web/public/samples`` and are served with the built web app.
"""

from pathlib import Path

import openpyxl
import pymupdf
from docx import Document
from docx.shared import Pt
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PptPt

OUT = Path(__file__).resolve().parent.parent / "public" / "samples"


def deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "第三季度销售汇报"
    title.placeholders[1].text = "示例公司 · 华南区销售部"
    slides = [
        (
            "区域业绩概览",
            [
                "华南区销售额同比增长百分之十二",
                "新客户数量达到历史最高水平",
                "线上渠道贡献了三成收入",
            ],
        ),
        (
            "下季度计划",
            ["扩大华东区的经销商网络", "推出两款新的节能产品", "加强售后服务团队的培训"],
        ),
    ]
    for heading, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for line in bullets[1:]:
            body.add_paragraph().text = line
        for paragraph in body.paragraphs:
            for run in paragraph.runs:
                run.font.size = PptPt(24)
    prs.save(str(path))


def contract(path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("業務委託契約書", level=1)
    doc.add_paragraph(
        "サンプル株式会社（以下「甲」という）と見本合同会社（以下「乙」という）は、"  # noqa: RUF001
        "次のとおり業務委託契約を締結する。"
    )
    doc.add_heading("第1条 委託業務", level=2)
    doc.add_paragraph("甲は乙に対し、社内文書の翻訳および校正の業務を委託し、乙はこれを受託する。")
    doc.add_heading("第2条 報酬及び支払", level=2)
    doc.add_paragraph("甲は乙に対し、毎月末日までに当月分の報酬を銀行振込により支払うものとする。")
    doc.add_paragraph("振込手数料は甲の負担とする。")
    doc.save(str(path))


def price_list(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.worksheets[0]
    ws.title = "Precios"
    ws.append(["Lista de precios de repuestos 2026"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Referencia", "Descripción", "Precio (EUR)", "Disponibilidad"])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F9C82")
    rows = [
        ("R-100", "Filtro de aire para compresor", 24.5, "En existencia"),
        ("R-205", "Correa de transmisión reforzada", 38.0, "Bajo pedido"),
        ("R-310", "Juego de juntas de goma", 12.75, "En existencia"),
        ("R-415", "Sensor de temperatura digital", 56.9, "Pocas unidades"),
    ]
    for row in rows:
        ws.append(row)
    for column, width in zip("ABCD", (14, 36, 14, 18), strict=True):
        ws.column_dimensions[column].width = width
    wb.save(path)


def manual(path: Path) -> None:
    doc = pymupdf.open()
    pages = [
        (
            "Manual de usuario",
            [
                "Gracias por elegir nuestro purificador de aire de ejemplo.",
                "Lea este manual antes de usar el aparato por primera vez.",
                "Conserve el manual para futuras consultas.",
            ],
        ),
        (
            "Instalación",
            [
                "Coloque el aparato sobre una superficie plana y estable.",
                "Deje al menos treinta centímetros libres a cada lado.",
                "Conecte el cable a una toma de corriente con toma de tierra.",
            ],
        ),
        (
            "Solución de problemas",
            [
                "Si el aparato no enciende, compruebe el cable de alimentación.",
                "Si hace ruido, limpie el filtro con un paño seco.",
                "Si el problema continúa, contacte con el servicio técnico.",
            ],
        ),
    ]
    for heading, lines in pages:
        page = doc.new_page(width=595, height=842)
        page.insert_text((72, 96), heading, fontsize=24, fontname="hebo")
        y = 150.0
        for line in lines:
            page.insert_text((72, y), line, fontsize=12, fontname="helv")
            y += 26
    doc.set_metadata({"title": "Manual de usuario (ejemplo)", "author": "Lenny samples"})
    doc.save(str(path), garbage=3, deflate=True)
    doc.close()


def notes(path: Path) -> None:
    path.write_text(
        "Weekly project meeting - sample notes\n"
        "\n"
        "Attendees: design, engineering and support teams.\n"
        "The new onboarding flow is ready for internal testing next Monday.\n"
        "Support asked for clearer error messages when an upload fails.\n"
        "Action: engineering will share a short demo before Friday.\n",
        encoding="utf-8",
    )


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    deck(OUT / "Q3 销售汇报.pptx")
    contract(OUT / "業務委託契約書_2026.docx")
    price_list(OUT / "Lista_de_precios_2026.xlsx")
    manual(OUT / "manual_de_usuario.pdf")
    notes(OUT / "meeting-notes.txt")
    for item in sorted(OUT.iterdir()):
        print(ascii(item.name), item.stat().st_size)


if __name__ == "__main__":
    main()
