#!/usr/bin/env python3

import argparse
import os
import torch

from docx import Document
from pptx import Presentation

from transformers import M2M100ForConditionalGeneration
from tokenization_small100 import SMALL100Tokenizer

# Define configuration
# Resolve the model directory relative to this script's location so the
# script works regardless of the current working directory.
MODEL_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "small100")  # Path to local model folder
SOURCE_LANGUAGE = "zh"               # Source language code
TARGET_LANGUAGE = "en"               # Target language code

print("Loading model and tokenizer from local folder...")
# Load tokenizer and model from local path
tokenizer = SMALL100Tokenizer.from_pretrained(MODEL_DIR)
model = M2M100ForConditionalGeneration.from_pretrained(MODEL_DIR)

# Device detection and model placement
device = "cuda" if torch.cuda.is_available() else "cpu"
model.to(device)
print(f"Using device: {device}")

# Configure source and target languages
tokenizer.src_lang = SOURCE_LANGUAGE
tokenizer.tgt_lang = TARGET_LANGUAGE

print(f"Translating from '{SOURCE_LANGUAGE}' to '{TARGET_LANGUAGE}'...")

def translate_texts(texts: list[str]) -> list[str]:
    """
    Translate a batch of texts from SOURCE_LANG to TARGET_LANG using a local M2M100/SMALL-100 model.
    Returns list of translated texts in same order as input.
    """
    if not texts:
        return []
    
    original_texts = texts
    texts = [t.strip() for t in texts]
    
    # Filter out empty texts but track their indices
    non_empty_indices = [i for i, t in enumerate(texts) if t]
    non_empty_texts = [texts[i] for i in non_empty_indices]
    
    if not non_empty_texts:
        return original_texts

    try:
        # Batch tokenize with padding and truncation
        encoded_inputs = tokenizer.batch_encode_plus(
            non_empty_texts,
            return_tensors="pt",
            padding=True,
            truncation=True,
            max_length=512
        )
        
        # Move to device
        encoded_inputs = {k: v.to(device) for k, v in encoded_inputs.items()}

        # Generate translations with optimized settings
        with torch.inference_mode():
            generated_tokens = model.generate(
                **encoded_inputs,
                use_cache=True,
                num_beams=1,
                do_sample=False,
                max_new_tokens=100
            )

        # Decode tokens back into readable strings
        translated_texts = tokenizer.batch_decode(generated_tokens, skip_special_tokens=True)

        # Map translations back to original positions
        results = list(original_texts)
        for idx, translated in zip(non_empty_indices, translated_texts):
            results[idx] = translated

        # Batch progress logging (safe for Windows console)
        print("-" * 60)
        for src, tgt in zip(non_empty_texts, translated_texts):
            # Encode to ASCII with replacement for console safety
            src_safe = src.encode('ascii', 'replace').decode('ascii')
            tgt_safe = tgt.encode('ascii', 'replace').decode('ascii')
            print(f"SRC: {src_safe}")
            print(f"TGT: {tgt_safe}")

        return results

    except Exception as ex:
        print(f"Batch translation failed: {ex}")
        return original_texts

def process_shape(shape):
    """
    Process textbox, table, placeholder, and group shapes while preserving formatting.
    """
    # Text boxes and placeholders
    if hasattr(shape, "text_frame") and shape.has_text_frame:
        process_text_frame(shape.text_frame)

    # Tables
    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            for cell in row.cells:
                if cell.text_frame:
                    process_text_frame(cell.text_frame)

    # Grouped shapes
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for subshape in shape.shapes:
            process_shape(subshape)

def process_text_frame(text_frame):
    """
    Translates text inside a text frame run-by-run to keep colors and formatting.
    Collects all runs first, batch translates, then maps results back.
    """
    # Collect all runs with text
    runs_to_translate = []
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                runs_to_translate.append(run)
    
    if not runs_to_translate:
        return
    
    # Extract texts preserving original whitespace
    texts = [run.text for run in runs_to_translate]
    
    # Batch translate
    translations = translate_texts(texts)
    
    # Map translations back to runs
    for run, translated in zip(runs_to_translate, translations):
        run.text = translated


def process_paragraph(paragraph):
    """
    Translates text inside a docx paragraph run-by-run to preserve formatting.
    Collects all runs with text, batch translates, then maps results back.
    """
    # Collect all runs with text
    runs_to_translate = []
    for run in paragraph.runs:
        if run.text.strip():
            runs_to_translate.append(run)

    if not runs_to_translate:
        return

    # Extract texts preserving original whitespace
    texts = [run.text for run in runs_to_translate]

    # Batch translate
    translations = translate_texts(texts)

    # Map translations back to runs
    for run, translated in zip(runs_to_translate, translations):
        run.text = translated


def translate_document(input_file, output_file):
    """
    Load a DOCX document, translate all paragraphs and table cells,
    and save the translated document.
    """
    doc = Document(input_file)

    # Process all paragraphs in the document body
    for paragraph in doc.paragraphs:
        process_paragraph(paragraph)

    # Process all tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)

    doc.save(output_file)

    print("\nTranslation complete.")
    print(f"Saved: {output_file}")


def translate_presentation(input_file, output_file):

    prs = Presentation(input_file)

    slide_count = len(prs.slides)

    for slide_num, slide in enumerate(prs.slides, start=1):

        print(f"\nProcessing slide {slide_num}/{slide_count}")

        for shape in slide.shapes:
            process_shape(shape)

    prs.save(output_file)

    print("\nTranslation complete.")
    print(f"Saved: {output_file}")


def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "input",
        help="Input file (.pptx or .docx)"
    )

    parser.add_argument(
        "output",
        help="Output translated file (.pptx or .docx)"
    )

    args = parser.parse_args()

    # Determine file type by extension
    input_ext = os.path.splitext(args.input)[1].lower()
    output_ext = os.path.splitext(args.output)[1].lower()

    if input_ext != output_ext:
        print(f"Error: Input and output file extensions must match (got {input_ext} and {output_ext})")
        return

    if input_ext == ".pptx":
        translate_presentation(args.input, args.output)
    elif input_ext == ".docx":
        translate_document(args.input, args.output)
    else:
        print(f"Error: Unsupported file extension '{input_ext}'. Supported: .pptx, .docx")
        return

if __name__ == "__main__":
    main()